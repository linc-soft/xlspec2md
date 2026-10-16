# -*- coding: utf-8 -*-
"""画面イメージ解析: 縦に積まれた複数画面を切り分け、結合レンジをピクセル座標付きコンポーネントに変換。

コンテナ階層（罫線矩形検出ベース、罫線が検出できないシートは従来の inset 推定にフォールバック）:
  Form（画面外枠 C11:DZ156、タイトルバー25pxを含む）
    └─ GroupBox（⊟ ラベル + 罫線矩形。y=ラベルブロック頂点、h=罫線枠高）
         └─ コンポーネント（rx=x-枠左、ry=y-ラベルブロック底）
    └─ Panel（ラベル無し罫線矩形、例: ボタン領域）
    └─ その他（ステータスバー等。ry はタイトル25pxを除いた値）
"""
from __future__ import annotations

import re

from openpyxl.utils import get_column_letter

from ..borders import BorderRect, detect_rectangles, strictly_inside
from ..models import Component, Screen
from ..wbutil import SheetView, col_px, row_px

FKEY_RE = re.compile(r"\((F\d{1,2}|ESC)\)")
GRID_ROW_RE = re.compile(r"^\d+\s*▶")


def classify(value) -> str:
    if value is None:
        return "block"
    s = str(value).strip()
    if not s:
        return "block"
    if s.startswith("■"):
        return "title"
    if s.startswith("⊟"):
        return "group"
    if FKEY_RE.search(s):
        return "button"
    if GRID_ROW_RE.match(s):
        return "grid-row"
    if s in ("▼", "▽", "▲", "△"):
        return "sort-or-dropdown"
    if s.startswith(("☐", "□")) and len(s) > 1:
        return "checkbox"
    return "text"


def parse_gajou(sv: SheetView, cfg: dict, warnings: list[str]) -> tuple[list[Screen], str, list[tuple[str, int, int]]]:
    """returns (screens, 機能名, annotations)"""
    # メタ: 機能名（"機能名"ラベルセルの右隣の結合レンジ値）
    func_name = ""
    hit = sv.find("機能名", max_row=cfg["header_end_row"])
    if hit:
        r, c = hit
        rng = sv.merge_of(r, c)
        next_col = (rng.max_col + 1) if rng else c + 1
        v = sv.value(r, next_col)
        if v:
            func_name = str(v).strip()

    # ---- 画面境界の検出 ----
    first_row = cfg["first_screen_row"]
    marker_col = cfg["screen_marker"]["column"]
    patterns = cfg["screen_marker"]["patterns"]
    min_components = cfg.get("min_screen_components", 4)
    last_content_row = max((r for r, _, _ in sv.iter_nonempty(min_row=first_row)), default=first_row)

    candidates = [(r, str(sv.value(r, marker_col)).strip())
                  for r in range(first_row + 1, sv.ws.max_row + 1)
                  if sv.value(r, marker_col) and sv.is_top_left(r, marker_col)
                  and any(p in str(sv.value(r, marker_col)) for p in patterns)]

    def count_merges(r0: int, r1: int) -> int:
        return sum(1 for rr in range(r0, r1 + 1) for cc in range(1, sv.ws.max_column + 1)
                   if sv.is_top_left(rr, cc) and sv.merge_of(rr, cc) is not None)

    # 候補ブロックを検査: コンポーネント数と密度の両方が閾値に達するものだけ画面とみなす。
    # 密度が低いブロック（例: 「メニュー」凡例）は注釈ブロックに格下げする。
    bounds: list[tuple[str, int]] = [("", first_row)]
    annotations: list[tuple[str, int, int]] = []  # (名前, 開始行, 終了行)
    cand_rows = [r for r, _ in candidates]
    following = cand_rows[1:] + [last_content_row + 1]
    min_density = cfg.get("min_screen_density", 0.15)
    for (r, name), nxt in zip(candidates, following):
        end = nxt - 1
        n = count_merges(r, end)
        density = n / max(1, end - r + 1)
        if n >= min_components and density >= min_density:
            bounds.append((name, r))
        else:
            annotations.append((name, r, end))
    ann_ranges = [(r0, r1) for _, r0, r1 in annotations]

    # 画面キャンバス（C11:DZ156 規約）。単画面シートの下端・全画面の幅の決定に使う。
    area = cfg.get("screen_area") or {}
    canvas_end_row = area.get("end_row")
    from openpyxl.utils import column_index_from_string
    canvas_end_col = column_index_from_string(area["end_col"]) if area.get("end_col") else None

    # 注釈ブロックを除いた内容のある最終行（単画面シートの下端決定用）
    def _not_ann(r: int) -> bool:
        return not any(r0 <= r <= r1 for r0, r1 in ann_ranges)

    last_content_excl = max((r for r, _, _ in sv.iter_nonempty(min_row=first_row) if _not_ann(r)),
                            default=first_row)

    # 各画面の終了行 = 次境界-1、最後は内容のある最終行
    screens: list[Screen] = []
    legend = set(cfg["legend_values"])

    for i, (marker_name, start) in enumerate(bounds):
        if i + 1 < len(bounds):
            end = bounds[i + 1][1] - 1
        else:
            # 単画面（末尾）シート: 注釈を除いた内容末行とキャンバス下端の大きい方
            end = max(last_content_excl, canvas_end_row or last_content_excl)
        name = func_name if i == 0 else marker_name
        # キャンバス右端が定義されていれば幅はそこから決める（シートのゴミ列に惑わされない）
        if canvas_end_col:
            width_px = (canvas_end_col - 3 + 1) * 10
        else:
            width_px = col_px(sv.ws.max_column + 1)
        screen = Screen(name=name, start_row=start, end_row=end,
                        width_px=width_px, height_px=row_px(end + 1, start))
        # ---- コンポーネント抽出 ----
        seen_ranges: set[str] = set()
        for r in range(start, end + 1):
            if any(r0 <= r <= r1 for r0, r1 in ann_ranges):
                continue  # 注釈ブロックは画面コンポーネントに含めない
            for c in range(1, sv.ws.max_column + 1):
                if not sv.is_top_left(r, c):
                    continue
                rng = sv.merge_of(r, c)
                if rng is None:
                    continue  # 単セル（罫線の名残）はコンポーネント化しない
                key = rng.coord
                if key in seen_ranges:
                    continue
                seen_ranges.add(key)
                v = sv.value(r, c)
                vs = str(v).strip() if v is not None else ""
                if vs in legend and r < start + 8:
                    continue  # 凡例
                comp = Component(
                    name=vs, type=classify(v),
                    x=col_px(rng.min_col), y=row_px(rng.min_row, start),
                    w=(rng.max_col - rng.min_col + 1) * 10,
                    h=(rng.max_row - rng.min_row + 1) * 5,
                    range=sv.range_str(rng), source=sv.name,
                )
                screen.components.append(comp)
        if screen.components:
            _assign_hierarchy(sv, screen, cfg, start, canvas_end_col, warnings)
            screens.append(screen)
        else:
            warnings.append(f"画面「{name}」(row {start}-{end}) にコンポーネントがありません")
    if not screens:
        warnings.append("画面イメージに画面が検出されませんでした")
    return screens, func_name, annotations


def _assign_hierarchy(sv: SheetView, screen: Screen, cfg: dict, start: int,
                      canvas_end_col: int | None, warnings: list[str]) -> None:
    """罫線矩形検出に基づくコンテナ階層構築（Form/GroupBox/Panel と相対座標）。

    罫線矩形から GroupBox を 1 つも特定できないシートは、従来の inset 推定
    （_assign_containers）にフォールバックする。"""
    comps = screen.components
    bb = cfg.get("border_box") or {}
    side = float(bb.get("side_coverage", 0.8))
    end_col = canvas_end_col or sv.ws.max_column
    rects = detect_rectangles(sv.ws, start, screen.end_row + 1, 3, end_col + 1,
                              side_coverage=side,
                              min_w_cols=int(bb.get("min_w_cols", 6)),
                              min_h_rows=int(bb.get("min_h_rows", 4)))

    groups = sorted((c for c in comps if c.type == "group"), key=lambda c: (c.y, c.x))
    titles = [c for c in comps if c.type == "title"]
    title_h = max((t.y + t.h for t in titles), default=0)
    others = [c for c in comps if c.type not in ("title", "group")]

    def cy_row(c: Component) -> int:
        return start + (c.y + c.h // 2) // 5

    def in_rect(c: Component, r: BorderRect) -> bool:
        cx, cy = c.x + c.w / 2, c.y + c.h / 2
        return (col_px(r.c1) <= cx <= col_px(r.c2)
                and row_px(r.r1, start) <= cy <= row_px(r.r2, start))

    # ---- GroupBox: ⊟ ラベルブロック → 罫線矩形マッチング ----
    # セクション = ラベルブロックの下から「次の ⊟ ラベル直上」まで。
    # 最終ラベルは「5行以上の空白」で区切る（ボタン領域・ステータスバーを
    # グループに呑み込ませないため）。
    pad = int(bb.get("label_row_pad", 2))
    matched: dict[BorderRect, tuple[Component, int]] = {}  # rect -> (group_comp, strip_h)
    for i, g in enumerate(groups):
        lt = start + g.y // 5
        lb = lt + g.h // 5 - 1
        if i + 1 < len(groups):
            sec_end = start + groups[i + 1].y // 5
        else:
            busy = {cy_row(c) for c in others}
            sec_end = screen.end_row + 1
            run = 0
            for r in range(lb + 1, screen.end_row + 1):
                run = run + 1 if r not in busy else 0
                if run >= 5:
                    sec_end = r - 4
                    break
        required = [c for c in others if lb < cy_row(c) < sec_end]
        lx1, lx2 = g.x, g.x + g.w
        cands = [r for r in rects
                 if lt - pad <= r.r1 <= lb + pad
                 and col_px(r.c1) <= lx1 and lx2 <= col_px(r.c2)
                 and all(in_rect(c, r) for c in required)]
        if not cands or not required:
            warnings.append(
                f"GroupBox「{g.name}」の罫線枠が特定できませんでした（ラベルブロックのまま配置します）")
            continue
        box = min(cands, key=lambda r: r.area)
        strip = g.h                       # ラベルブロック高（上書き前に保存）
        bx, by = col_px(box.c1), row_px(box.r1, start)
        g.label_offset = by - g.y         # 罫線枠上端 = y + label_offset（通常 10）
        g.x, g.w, g.h = bx, (box.c2 - box.c1) * 10, (box.r2 - box.r1) * 5
        matched[box] = (g, strip)

    if not matched:
        # 罫線が検出できないシート: 従来の inset 推定にフォールバック
        inset_x = cfg.get("group_inset_col", 1) * 10
        inset_y = cfg.get("group_inset_row", 1) * 5
        _assign_containers(screen, inset_x, inset_y)
        return

    # ---- Form コンポーネント（画面外枠。キャンバス規約 C11:DZ156 から合成）----
    form = Component(name=screen.name, type="form", x=0, y=0,
                     w=screen.width_px, h=screen.height_px,
                     range=f"{sv.name}!{get_column_letter(3)}{start}:"
                           f"{get_column_letter(end_col)}{screen.end_row}",
                     source=sv.name)

    # ---- Panel: ラベル無し罫線矩形（ボタン領域など）----
    # 重なり合う候補はクラスタにまとめ、各クラスタから「外枠列（C/EA）に触れない
    # 最大面積」の矩形を採用する（ボタン同士の共有罫線が細かい候補を大量に生成するため。
    # 外枠列との重なり候補は面積が 1 列分だけ大きくなることが多く、これを除外する）。
    # 内包判定は ±1 列/行の余裕を持たせる（罫線が 1 セルずれて枠の外に突き出すため）。
    def loosely_inside(a: BorderRect, b: BorderRect, tol: int = 1) -> bool:
        return (b.r1 - tol <= a.r1 and a.r2 <= b.r2 + tol
                and b.c1 - tol <= a.c1 and a.c2 <= b.c2 + tol)

    kept = [r for r in rects if r not in matched
            and not any(loosely_inside(r, m) for m in matched)]
    cands: list[BorderRect] = []
    for r in kept:
        if any(loosely_inside(m, r) for m in matched):
            continue  # GroupBox を内包する誤矩形（外枠との重なり等）は除外
        cs = [c for c in others if in_rect(c, r)]
        if len(cs) >= 2:
            cands.append(r)

    def intersects(a: BorderRect, b: BorderRect) -> bool:
        return not (a.r2 <= b.r1 or b.r2 <= a.r1 or a.c2 <= b.c1 or b.c2 <= a.c1)

    clusters: list[list[BorderRect]] = []
    for r in cands:
        hit = [cl for cl in clusters if any(intersects(r, x) for x in cl)]
        if not hit:
            clusters.append([r])
        else:
            hit[0].append(r)
            for extra in hit[1:]:
                hit[0].extend(extra)
                clusters.remove(extra)

    def in_canvas(r: BorderRect) -> bool:
        return r.c1 > 3 and r.c2 <= end_col

    panels: list[tuple[BorderRect, Component]] = []
    for cl in clusters:
        preferred = [r for r in cl if in_canvas(r)] or cl
        r = max(preferred, key=lambda x: x.area)
        p = Component(name=f"□ {sv.coord(r.r1, r.c1)}:{sv.coord(r.r2, r.c2)}",
                      type="panel", x=col_px(r.c1), y=row_px(r.r1, start),
                      w=(r.c2 - r.c1) * 10, h=(r.r2 - r.r1) * 5,
                      range=f"{sv.name}!{sv.coord(r.r1, r.c1)}:{sv.coord(r.r2, r.c2)}",
                      source=sv.name)
        panels.append((r, p))

    # ---- コンテナ割当（面積最小の内包コンテナ = 最も内側）----
    containers: list[tuple[int, str, tuple[int, int, int, int], tuple[int, int]]] = []
    for rect, (g, strip) in matched.items():
        bx, by = col_px(rect.c1), row_px(rect.r1, start)
        bw, bh = (rect.c2 - rect.c1) * 10, (rect.r2 - rect.r1) * 5
        g.container = form.name
        g.rx, g.ry = bx, g.y - title_h
        # 子の相対座標原点: x=罫線枠左、y=ラベルブロック底（strip 分下）
        containers.append((bw * bh, g.name, (bx, by, bw, bh), (bx, g.y + strip)))
    panel_ids = {id(p) for _, p in panels}
    for r, p in panels:
        p.container = form.name
        p.rx, p.ry = p.x, p.y - title_h
        containers.append((p.w * p.h, p.name, (p.x, p.y, p.w, p.h), (p.x, p.y)))
    containers.append((screen.width_px * screen.height_px, form.name,
                       (0, 0, screen.width_px, screen.height_px), (0, title_h)))

    for c in comps:
        if c.type in ("title", "group", "form") or id(c) in panel_ids:
            continue
        if title_h and c.y + c.h <= title_h:
            continue  # タイトルバーは Form の chrome（コンテナ割当の対象外）
        cx, cy = c.x + c.w / 2, c.y + c.h / 2
        best: tuple[int, str, tuple[int, int]] | None = None
        for area, name, (rx_, ry_, rw_, rh_), (ox, oy) in containers:
            if rx_ <= cx <= rx_ + rw_ and ry_ <= cy <= ry_ + rh_:
                if best is None or area < best[0]:
                    best = (area, name, (ox, oy))
        if best is None:
            continue
        _, name, (ox, oy) = best
        c.container = name
        c.rx, c.ry = c.x - ox, c.y - oy

    for g in groups:
        if g.container == "":
            g.container = form.name
            g.rx, g.ry = g.x, g.y - title_h

    # Form を先頭に、パネルも含めて y,x 順に並べ替え
    screen.components = [form] + sorted(
        [c for c in comps if c is not form] + [p for _, p in panels],
        key=lambda c: (c.y, c.x, c.x + c.w))


def _assign_containers(screen: Screen, inset_x: int, inset_y: int) -> None:
    """group 型コンポーネントをコンテナとみなし、内包する他コンポーネントに
    コンテナ名とコンテナ内相対座標を設定する（罫線非検出シート用の従来処理）。
    グループ罫線の原点 = ⊟ ラベルの座標 - inset（罫線はラベルより外側にあるため）。"""
    groups = [(g.x - inset_x, g.y - inset_y, g) for g in screen.components if g.type == "group"]
    for comp in screen.components:
        if comp.type == "group":
            continue
        # 原点が自分より上左にあるグループのうち、最も近い（最も内側の）ものを親とする
        candidates = [(gx, gy, g) for gx, gy, g in groups
                      if comp.x >= gx and comp.y >= gy]
        if not candidates:
            continue
        gx, gy, g = max(candidates, key=lambda t: (t[1], t[0]))
        comp.container = g.name
        comp.rx = comp.x - gx
        comp.ry = comp.y - gy
