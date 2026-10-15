# -*- coding: utf-8 -*-
"""ワークブック読み込み＆セル/結合レンジ共通操作。"""
from __future__ import annotations

import re
from typing import Iterator, Optional

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.worksheet.cell_range import CellRange


def load_wb(path: str):
    return openpyxl.load_workbook(path, data_only=True)


class SheetView:
    """結合セルを意識したシートビュー。value() は結合レンジの左上値を返す。"""

    def __init__(self, ws: Worksheet):
        self.ws = ws
        self.name = ws.title
        # (row, col) -> CellRange（そのセルを覆う結合レンジ）
        self.merge_at: dict[tuple[int, int], CellRange] = {}
        for rng in ws.merged_cells.ranges:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    self.merge_at[(r, c)] = rng

    # ---- 値アクセス ----
    def raw(self, row: int, col: int):
        return self.ws.cell(row=row, column=col).value

    def value(self, row: int, col: int):
        """結合レンジに属すれば左上セルの値、さもなくば自セルの値。"""
        rng = self.merge_at.get((row, col))
        if rng is not None:
            return self.ws.cell(row=rng.min_row, column=rng.min_col).value
        return self.ws.cell(row=row, column=col).value

    def merge_of(self, row: int, col: int) -> Optional[CellRange]:
        return self.merge_at.get((row, col))

    def is_top_left(self, row: int, col: int) -> bool:
        rng = self.merge_at.get((row, col))
        return rng is None or (rng.min_row == row and rng.min_col == col)

    # ---- 走査 ----
    def iter_nonempty(self, min_row: int = 1, max_row: Optional[int] = None,
                      min_col: int = 1, max_col: Optional[int] = None) -> Iterator[tuple[int, int, object]]:
        """値のあるセルを (row, col, value) で返す。結合レンジは左上のみ。"""
        max_row = max_row or self.ws.max_row
        max_col = max_col or self.ws.max_column
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                if not self.is_top_left(r, c):
                    continue
                v = self.value(r, c)
                if v is not None and str(v).strip():
                    yield r, c, v

    def find(self, keyword: str, min_row: int = 1, max_row: Optional[int] = None,
             min_col: int = 1, max_col: Optional[int] = None) -> Optional[tuple[int, int]]:
        """keyword を含む先頭セル（左上基準）を返す。"""
        for r, c, v in self.iter_nonempty(min_row, max_row, min_col, max_col):
            if keyword in str(v):
                return r, c
        return None

    def coord(self, row: int, col: int) -> str:
        return f"{get_column_letter(col)}{row}"

    def range_str(self, rng: CellRange) -> str:
        return f"{self.name}!{rng.coord}"


# ---- 幾何計算（1列=10px、1行=5px、列C=画面左端=0px）----
COL_ORIGIN = 3   # C 列
PX_PER_COL = 10
PX_PER_ROW = 5

FULLWIDTH_DIGITS = str.maketrans("０１２３４５６７８９．", "0123456789.")


def col_px(col: int) -> int:
    return (col - COL_ORIGIN) * PX_PER_COL


def row_px(row: int, origin_row: int) -> int:
    return (row - origin_row) * PX_PER_ROW


def norm_fullwidth(s: str) -> str:
    return s.translate(FULLWIDTH_DIGITS)
